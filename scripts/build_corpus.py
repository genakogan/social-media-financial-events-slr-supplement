#!/usr/bin/env python3
"""Build corpus register (Supplementary Table S1) for the 126-study corpus, in the style of corpus_v6."""
import csv, re, sys, json, math
from collections import OrderedDict
sys.path.insert(0, '.')
from extract_citations import parse_bib
from pylatexenc.latex2text import LatexNodes2Text
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = 'Supplementary_Table_S1_Corpus_Register.xlsx'
L2T = LatexNodes2Text()
tx = lambda s: re.sub(r'\s+', ' ', L2T.latex_to_text(s or '')).strip()
bib = {e['key']: e for e in parse_bib('surveybib_corrected.bib')}
occ = list(csv.DictReader(open('out_v2/citation_occurrences.csv', encoding='utf-8')))  # citation inventory of the revised manuscript

# ---------------- corpus definition by PRISMA box (order = manuscript order)
T_SM, T_FIN, T_X = 'Social media analysis', 'Financial market events', 'Cross-Domain Events and Causality'
BOXES = OrderedDict([
    ('Opinion in Social Networks', (T_SM, '3.2', 'How Opinion Is Defined',
        'vandenbos2007apa musch2003psychology zhang4450056integrated degroot1974reaching hegselmann2002opinion friedkin1990social acemoglu2011opinion gaillard2021countering kietzmann2011social')),
    ('Opinion Manipulation: Definition and Typologies', (T_SM, '3.3', 'Opinion Manipulation and Typology',
        'ienca2023artificial cayuela2025influence chang2022systematic pastor2025influence pal2025media kumarsocial')),
    ('Class 1 - Mathematical and System-Dynamics Modelling', (T_SM, '3.4', 'Opinion Detection and Modeling (C1)',
        'paul2023mathematical ge2024novel evangelatos2025modeling zhou2020two butts2023mathematical bhalla2025local')),
    ('Class 2 - Machine & Deep Learning Content Detection', (T_SM, '3.4', 'Opinion Detection and Modeling (C2)',
        'mridha2021comprehensive saadi2025enhancing huang2025hybrid tufchi2024transvae sharma2025novel ozcelik2025detecting yang2025llm mahmud2026detecting')),
    ('Class 3 - Hybrid Human-AI Annotation and Verification', (T_SM, '3.4', 'Opinion Detection and Modeling (C3)',
        'sahitaj2025hybrid fan2026skeptik liu2024exploring hrckova2025autonomation cui2025toward lee2025countering')),
    ('Class 4 - Network Structure and Propagation Analysis', (T_SM, '3.4', 'Opinion Detection and Modeling (C4)',
        'nobre2022hierarchical resende2019analyzing chen2022social loru2026compression bellina2025language impicciche2025comparing wang2026rumor fawzi2024pinocchio')),
    ('Class 5 - Behavioural and Psychological Modelling', (T_SM, '3.4', 'Opinion Detection and Modeling (C5)',
        'escola2023falls smith2024unraveling lutz2024linguistic sakib2025understanding amieur2025comparative budak2024misunderstanding khan2024targets govers2025feeds')),
    ('Class 6 - Intervention and Inoculation Strategies', (T_SM, '3.4', 'Opinion Detection and Modeling (C6)',
        'roozenbeek2022technique wang2025limited tang2025breaking chuai2024did jeon2024hearhere yu2025breaking su2025emotionally shahid2025one')),
    ('F1 - Classical event-study econometrics', (T_FIN, '3.6', 'Methodological Approaches in Financial Event Research (F1)',
        'mackinlay1997event el2023event corrado1989nonparametric boehmer1991event pynnonen2022non nguyen2026permutation kothari2007econometrics goldsmith2025causal')),
    ('F2 - Volatility, risk, and structural change', (T_FIN, '3.6', 'Methodological Approaches in Financial Event Research (F2)',
        'engle1982autoregressive bollerslev1986generalized zhu2025volatility bai2003computation hasanov2024structural sun2026inference mugrabi2026financial')),
    ('F3 - High-frequency and intraday market response', (T_FIN, '3.6', 'Methodological Approaches in Financial Event Research (F3)',
        'anderson2003micro hasbrouck1991measuring bacry2015hawkes nystrom2022hawkes mucciante2024estimation jusselin2020no gross2011machines yang2025price behrendt2018twitter')),
    ('F4 - Textual and attention-augmented asset pricing', (T_FIN, '3.6', 'Methodological Approaches in Financial Event Research (F4)',
        'tetlock2007giving loughran2011liability da2011search chen2014wisdom liu2023social kim2025news kraaijeveld2020predictive arnold2022attention')),
    ('F5 - Market manipulation and abusive trading detection', (T_FIN, '3.6', 'Methodological Approaches in Financial Event Research (F5)',
        'xu2019anatomy kogan2023social leuz2026falls li2025cryptocurrency la2023doge bradley2024place keasey2025impact roy2024unveiling')),
    ('Social Media signals', (T_X, '4.1.1', 'SM Events',
        'cordeiro2016online dou2012event becker2009event coviello2014detecting bayer2020social riva2016psychology')),
    ('Financial signals', (T_X, '4.1.2', 'Financial Events', 'campbell1997econometrics westaby2014psychology')),
    ('Cross-Domain Challenges', (T_X, '4.2', 'Cross-Domain Research',
        'wang2022causal yang2018applications ding2021quantifying arcuri2023does fan2024measuring')),
    ('Causal Identification', (T_X, '4.2.1', 'Causal Identification in Cross-Domain Research',
        'granger1969investigating sims1980macroeconomics schreiber2000measuring card1993minimum imbens2008regression abadie2010synthetic assaad2022survey gong2024causal chen2024discovery runge2019detecting sadeghi2025causal tong2022assessing seetharam2024headlines ghazouani2025virtual')),
])
# diagram values: (total incl. overlap, overlap)
DIAGRAM = {'Opinion in Social Networks': (9, 0), 'Opinion Manipulation: Definition and Typologies': (6, 0),
           'Class 1 - Mathematical and System-Dynamics Modelling': (6, 0), 'Class 2 - Machine & Deep Learning Content Detection': (8, 0),
           'Class 3 - Hybrid Human-AI Annotation and Verification': (6, 0), 'Class 4 - Network Structure and Propagation Analysis': (8, 0),
           'Class 5 - Behavioural and Psychological Modelling': (8, 0), 'Class 6 - Intervention and Inoculation Strategies': (8, 0),
           'F1 - Classical event-study econometrics': (8, 0), 'F2 - Volatility, risk, and structural change': (7, 0),
           'F3 - High-frequency and intraday market response': (9, 0), 'F4 - Textual and attention-augmented asset pricing': (8, 0),
           'F5 - Market manipulation and abusive trading detection': (8, 0),
           'Social Media signals': (6, 0), 'Financial signals': (4, 2), 'Cross-Domain Challenges': (6, 1), 'Causal Identification': (19, 5)}
XREF = {'dou2012event': ['Financial signals'], 'bayer2020social': ['Financial signals'],
        'xu2019anatomy': ['Cross-Domain Challenges'],
        'mackinlay1997event': ['Causal Identification'], 'campbell1997econometrics': ['Causal Identification'],
        'goldsmith2025causal': ['Causal Identification'], 'chuai2024did': ['Causal Identification'],
        'li2025cryptocurrency': ['Causal Identification']}
CLASSES = {'C1': 'Class 1', 'C2': 'Class 2', 'C3': 'Class 3', 'C4': 'Class 4', 'C5': 'Class 5', 'C6': 'Class 6'}
FOUNDATIONAL = set('''vandenbos2007apa musch2003psychology zhang4450056integrated degroot1974reaching hegselmann2002opinion friedkin1990social
acemoglu2011opinion kietzmann2011social mackinlay1997event corrado1989nonparametric boehmer1991event kothari2007econometrics el2023event
engle1982autoregressive bollerslev1986generalized bai2003computation anderson2003micro hasbrouck1991measuring bacry2015hawkes
tetlock2007giving loughran2011liability da2011search cordeiro2016online dou2012event becker2009event campbell1997econometrics
granger1969investigating sims1980macroeconomics schreiber2000measuring card1993minimum imbens2008regression abadie2010synthetic
assaad2022survey gong2024causal chen2024discovery runge2019detecting sadeghi2025causal'''.split())
RESTORED = {'tong2022assessing': '4.2.1', 'seetharam2024headlines': '4.2.1', 'ghazouani2025virtual': '4.2.1', 'wang2022causal': '4.2'}
FULLTEXT_CHECKED = {'seetharam2024headlines', 'ghazouani2025virtual', 'wang2022causal'}

SUPPORTING = OrderedDict([
    ('kumar2018false', ('1.1', 'Background and motivation (information ecosystem)')),
    ('armstrong2019organized', ('1.1', 'Historical motivation; possible citation-claim mismatch (chapter title concerns Roman army organisation)')),
    ('nam2025detecting', ('1.2', 'Background (algorithmic amplification)')),
    ('silva2025guidelines', ('1.2', 'Software-engineering motivation')),
    ('nazir2024architecting', ('1.2', 'Software-engineering motivation')),
    ('rahman2025secureai', ('1.2', 'Software-engineering motivation')),
    ('chen2025year', ('1.2', 'Motivation (experimental prototypes)')),
    ('Cortis2020SocialOpinionMining', ('1.4', 'Prior survey used to position the review')),
    ('Reveilhac2022SocialMediaSurveys', ('1.4', 'Prior survey used to position the review')),
    ('Abkenar2020BigDataSocialMedia', ('1.4', 'Prior survey used to position the review')),
    ('Rahman2022SystematicReviewBigDataSocialMedia', ('1.4', 'Prior survey used to position the review')),
    ('Long2025SocialMediaCapitalMarkets', ('1.4', 'Prior survey used to position the review')),
    ('Sundarasen2025TweetsToTrades', ('1.4', 'Prior survey used to position the review')),
    ('Doogan2023TopicModelsShortTextSocialMedia', ('1.4', 'Prior survey used to position the review')),
    ('kitchenham2004evidence', ('2', 'SLR methodology guideline')),
    ('stol2015holistic', ('2', 'SLR methodology guideline')),
    ('hinderks2020slr', ('2', 'SLR methodology guideline')),
    ('stapic2012performing', ('2', 'SLR methodology guideline')),
    ('page2021prisma', ('2.2', 'PRISMA 2020 reporting guideline')),
    ('hausner2012routine', ('2 (Reproducibility of the Review)', 'Method for objectively derived search strategies (known-item retrieval check); cited in the revision')),
])

def pubtype(e):
    t, f = e['type'], e['fields']
    j = (f.get('journal', '') + ' ' + f.get('note', '') + ' ' + f.get('howpublished', '')).lower()
    if t == 'book': return 'Book'
    if t == 'incollection': return 'Book chapter'
    if t == 'inproceedings': return 'Conference paper'
    if t in ('misc', 'techreport') or 'arxiv' in j or 'ssrn' in j: return 'Preprint / working paper'
    return 'Journal article'

def authors(k, e):
    a = e['fields'].get('author', '')
    parts = [tx(p) for p in re.split(r'\s+and\s+', a)] if a else []
    return '; '.join('et al.' if p.lower() == 'others' else p for p in parts)

def title(k, e):
    return tx(e['fields'].get('title', ''))

def doi(e):
    d = e['fields'].get('doi', '').strip()
    d = re.sub(r'^(https?://)?(dx\.)?doi\.org/|^doi:\s*', '', d)
    return d or None

def venue(e):
    f = e['fields']
    for k in ('journal', 'booktitle', 'publisher', 'institution', 'howpublished'):
        if f.get(k): return tx(f[k])
    return None

def year(e):
    y = e['fields'].get('year', '')
    return int(y) if y.isdigit() else y

def cited_sections(k):
    s = []
    for o in occ:
        if o['key'] == k:
            lab = 'Appendix A.1' if o['section_number'].startswith('A') else f"{o['section_number']} {o['section_title']}"
            if lab not in s: s.append(lab)
    if k in RESTORED and not any(x.startswith(RESTORED[k] + ' ') for x in s):   # restored citation not yet in the text
        lab = '4.2.1 Causal Identification in Cross-Domain Research' if RESTORED[k] == '4.2.1' else '4.2 Cross-Domain Research'
        s.append(lab)
    return '; '.join(s)

CORR = {}
for r in csv.DictReader(open('bib_corrections_log.csv', encoding='utf-8')):
    CORR.setdefault(r['bibkey'], []).append(f"{r['field']}: {r['reason']}")
NOTES = {
    'wang2026rumor': 'Volume/issue/pages in .bib look copied from another entry; verify.',
    'kumarsocial': 'Weak record: no volume, issue, pages or DOI; not indexed in OpenAlex.',
    'cayuela2025influence': 'arXiv preprint; not indexed in OpenAlex.',
    'pal2025media': 'Book chapter recorded as @article in .bib.',
    'riva2016psychology': 'Book chapter recorded as @article in .bib.',
    'sahitaj2025hybrid': 'Recorded as arXiv preprint; DOI is the published workshop version.',
    'huang2025hybrid': 'Recorded as arXiv preprint; DOI is the published IEEE conference version.',
    'cui2025toward': 'Recorded as arXiv preprint; DOI is the published ACM version.',
    'liu2024exploring': 'Recorded as arXiv preprint; DOI is the published ACM version.',
    'pastor2025influence': 'Cited as arXiv preprint; OpenAlex lists a published version (IEEE Internet Computing, DOI 10.1109/MIC.2026.3700710). Optional update; would change publication-type counts.',
    'paul2023mathematical': 'Cited with the arXiv DOI; OpenAlex lists the published chapter (DOI 10.1007/978-981-97-9194-1_20). Optional update.',
    'mackinlay1997event': 'DOI 10.1257/jel.35.1.13 is not registered in Crossref; verify or replace with the JSTOR stable URL.',
}
for k, v in CORR.items():
    NOTES[k] = ('Corrected in surveybib_corrected.bib (2026-10-19): ' + '; '.join(v)
                + (' ' + NOTES[k] if k in NOTES else ''))

# ---------------- metadata check (OpenAlex record of each study)
RES = {r['bibkey']: r for r in csv.DictReader(open('run6/corpus_resolved.csv', encoding='utf-8'))}
# ---------------- validated verification search (val/PROTOCOL.md; final outputs in val/run_final)
VAL_SPLIT = {r['bibkey']: r['set'] for r in csv.DictReader(open('val/split.csv', encoding='utf-8'))}
VAL_RET = {r['bibkey']: r for r in csv.DictReader(open('val/run_final/retrieval.csv', encoding='utf-8'))}
VAL_LOG = json.load(open('val/run_final/protocol_log.json', encoding='utf-8'))
VAL_TERMS = list(csv.DictReader(open('val/run_final/term_counts.csv', encoding='utf-8')))
V2 = json.load(open('val/run_v2/results_v2.json', encoding='utf-8'))
FIELDS = list(csv.DictReader(open('val/run_v2/fields_check_console.csv', encoding='utf-8')))
FINAL = json.load(open('val/run_final2/results_final.json', encoding='utf-8'))
TYPES_SUM = json.load(open('val/run_final2/types_summary.json', encoding='utf-8'))
assert len(FIELDS) == 82
assert set(VAL_SPLIT) == set(VAL_RET) and len(VAL_SPLIT) == 82

def val_retrieved(k):
    r = VAL_RET.get(k)
    if not r:
        return None
    if r['indexed'] != 'True':
        return 'not indexed in OpenAlex'
    return 'yes' if r['retrieved'] == 'True' else 'no'

def val_final(k):
    # final query (search_protocol): validation set only
    if VAL_SPLIT.get(k) != 'validation':
        return None
    if k in FINAL['validation_not_indexed']:
        return 'not indexed in OpenAlex'
    return 'no' if k in FINAL['validation_missed'] else 'yes'

ARXIV = ('10.48550/arxiv',)
# OpenAlex record reached through the wrong DOI is a journal item about the book, not the book itself
RES['campbell1997econometrics'] = dict(RES.get('campbell1997econometrics', {}), oa_id='')

def meta_check(k):
    if k == 'campbell1997econometrics':
        return 'Book, ISBN 9780691043012; the OpenAlex/Crossref record reached through the old DOI is a 1998 journal item about the book'
    r = RES.get(k, {})
    if not r.get('oa_id'):
        return 'Not indexed in OpenAlex'
    s_ = f"OpenAlex {r['oa_id']}"
    if k in CORR and any(c.startswith('doi') for c in CORR[k]):
        return s_ + ' (matched by title); DOI corrected in .bib'
    if r.get('cr_status') == 'ok':
        return s_ + '; Crossref DOI record checked'
    if (r.get('doi') or '').lower().startswith(ARXIV):
        return s_ + '; arXiv DOI (registered with DataCite, not Crossref)'
    if r.get('cr_status') == 'no DOI':
        return s_ + ' (matched by title); no DOI in .bib'
    return s_ + '; DOI not registered in Crossref'

rows = []
for box, (theme, sec, secname, keys) in BOXES.items():
    for k in keys.split():
        e = bib[k]
        cls = None
        m = re.match(r'(Class (\d)|F(\d))', box)
        if box.startswith('Class'): cls = 'C' + box[6]
        elif box.startswith('F') and box[1].isdigit(): cls = box[:2]
        apx = f'Table {cls}' if cls else None
        mt = meta_check(k) + ('; full text checked' if k in FULLTEXT_CHECKED else '')
        rows.append(OrderedDict([
            ('study_id', None), ('bibkey', k), ('authors', authors(k, e)), ('year', year(e)), ('title', title(k, e)),
            ('venue', venue(e)), ('publication_type', pubtype(e)), ('doi', doi(e)),
            ('theme', theme), ('prisma_box', box), ('xref_to', '; '.join(XREF.get(k, [])) or None),
            ('section_no', sec), ('section', secname), ('method_class', cls), ('appendix_table', apx),
            ('cited_in', cited_sections(k)),
            ('study_role', 'Foundational' if k in FOUNDATIONAL else 'Core'),
            ('revision_status', 'restored to text in revision' if k in RESTORED else None),
            ('metadata_check', mt), ('note', NOTES.get(k)),
            ('openalex_id', RES.get(k, {}).get('oa_id') or None),
            ('in_search_window', 'yes' if isinstance(year(e), int) and year(e) >= 2021 else 'no'),
            ('validation_set', VAL_SPLIT.get(k)), ('retrieved_by_final_query', val_final(k))]))
for i, r in enumerate(rows, 1):
    r['study_id'] = f'S{i:03d}'
assert len(rows) == 126 and len({r['bibkey'] for r in rows}) == 126, len(rows)

sup = []
for k, (sec, why) in SUPPORTING.items():
    e = bib[k]
    sup.append(OrderedDict([('bibkey', k), ('authors', authors(k, e)), ('year', year(e)), ('title', title(k, e)),
                            ('publication_type', pubtype(e)), ('doi', doi(e)), ('cited_in_section', sec),
                            ('reason not in corpus', why + '. Cited outside the sections counted in the PRISMA categories.')]))

# ---------------- workbook
AR = 'Arial'
HF = Font(name=AR, bold=True, color='FFFFFF', size=10); HFILL = PatternFill('solid', start_color='44546A')
BF = Font(name=AR, size=9); BLUE = Font(name=AR, size=9, color='0000FF'); BOLD = Font(name=AR, size=9, bold=True)
thin = Side(style='thin', color='BFBFBF'); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical='top')

def table(ws, headers, data, widths, start=1):
    for c, h in enumerate(headers, 1):
        x = ws.cell(row=start, column=c, value=h); x.font, x.fill, x.border = HF, HFILL, BD; x.alignment = WR
    for r, d in enumerate(data, start + 1):
        for c, h in enumerate(headers, 1):
            x = ws.cell(row=r, column=c, value=d[h] if isinstance(d, dict) else d[c - 1])
            x.font, x.border, x.alignment = BF, BD, WR
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=3 if len(headers) > 4 else 1)
    ws.auto_filter.ref = f'A{start}:{get_column_letter(len(headers))}{start + len(data)}'
    return start + len(data)

wb = Workbook()
ws = wb.active; ws.title = 'README'
readme = [
    ('Supplementary Table S1 - Register of the review corpus (126 studies)', Font(name=AR, bold=True, size=13)),
    ('Connecting Social Media Analysis and Financial Market Events: A Systematic Survey', BF),
    ('', BF),
    ('What this file is', BOLD),
    ('A study-level register of the 126 works included in the final synthesis. Each record is linked to its PRISMA category (Figure: PRISMA flow diagram), '
     'the manuscript section in which it is analysed and, where applicable, its methodological class (C1-C6, F1-F5) and appendix table.', BF),
    ('The register was compiled from the final manuscript and its bibliographic database on 18 October 2026 and synchronised with the revised manuscript on 19 October 2026 (all 146 cited keys: 126 in the corpus, 20 on not_in_corpus); bibliographic data are taken from surveybib_corrected.bib. Scripts and data files are described in the README.md of the accompanying repository.', BF),
    ('', BF),
    ('How the corpus was checked', BOLD),
    ('Every PRISMA category count is recomputed from the Data sheet (sheet PRISMA_check) and matches the flow diagram: social media analysis 59, '
     'financial market events 40, cross-domain events and causality 27 (total 126). Publication types are recomputed on the Counts sheet and match the '
     'corpus-composition table (103 journal articles, 7 conference papers, 3 books, 2 book chapters, 11 preprints/working papers).', BF),
    ('', BF),
    ('Metadata check', BOLD),
    ('On 19 October 2026 the bibliographic metadata of all 126 studies were checked against OpenAlex and their DOIs against Crossref (verify_corpus.py); '
     'corrections to the bibliographic database are listed on sheet bib_corrections.', BF),
    ('', BF),
    ('Retrospective known-item retrieval check (supplementary; sheets search_protocol, validation_result, validation_terms, protocol_history, openalex_fields)', BOLD),
    ('One protocol (search_protocol), following the objective approach of Hausner et al. (2012, Systematic Reviews 1:19). The 82 included studies published 2021-2026 were split at random '
     'within each flow-diagram category (seed 20261019) into a development set (44) and a validation set (38). Candidate terms were taken from the development titles only (words and '
     'two-word phrases in at least two titles) and from the keywords of Table 2; terms over-represented in the development titles relative to OpenAlex were kept, with the threshold chosen '
     'on the development set (R = 100), giving a title query of 27 terms (frozen, SHA-256 recorded). Restrictions: publication dates 2020-2026, ineligible document types excluded, '
     'OpenAlex subject fields of the development studies; no language restriction.', BF),
    ('Result (20 October 2026): 311,556 records (528,902 without the subject-field restriction on the same day; -41.1%); validation set 29 of 38 studies retrieved (76.3%, 95% CI 60.8-87.0%), '
     '29 of the 36 indexed in OpenAlex (80.6%, 95% CI 65.0-90.2%); development set 42 of 44. The subject-field restriction was chosen after the fields of the validation studies had been inspected, '
     'so it is not independently validated; it did not remove any of the known included studies retrieved by the original query. The final configuration is used for the reported filtered retrieval result, while the original pre-specified configuration and all intermediate runs are retained as part of the methodological audit trail. The original configuration, a language filter that lost two studies without an '
     'assigned language, and rejected phrase-based variants are on sheet protocol_history. The check measures retrieval of known included studies; because both samples come from the included corpus, '
     'it cannot establish the completeness or sensitivity of the original search. Per-study results: Data, columns validation_set and retrieved_by_final_query.', BF),
    ('', BF),
    ('Sheets', BOLD),
    ('Data - one row per included study. | not_in_corpus - ' + str(len(sup)) + ' works cited in the manuscript only for background or methodological guidance. | '
     'codebook - definition of every column and label. | class_definitions - definitions of classes C1-C6 and F1-F5 quoted from the manuscript, with the assignment procedure and counts (formulas). | Counts - publication types, themes and classes (formulas). | PRISMA_check - diagram vs register (formulas). | '
     'search_concepts_table2 - search concepts and Boolean strings as reported in the manuscript. | search_protocol - protocol of the known-item retrieval check. | '
     'validation_result - filters, frozen query, threshold, records, retrieval in the development and validation sets, per-study result. | validation_terms - candidate terms with OpenAlex counts and over-representation ratios. | '
     'protocol_history - methodological audit trail: all configurations run (original pre-specified, intermediate, dropped, rejected and final), with records and retrieval; records by OpenAlex field. | openalex_fields - OpenAlex field, subfield and primary topic of the 82 studies published 2021-2026. | '
     'bib_corrections - corrections applied to surveybib.bib. | decisions_log - decisions taken when compiling the register.', BF),
]
for i, (t, f) in enumerate(readme, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[i].height = max(15, 13 * (len(t) // 150 + 1))
ws.column_dimensions['A'].width = 160

ws = wb.create_sheet('Data')
H = list(rows[0].keys())
DEND = table(ws, H, rows, [8, 26, 34, 6, 50, 30, 16, 26, 22, 34, 20, 7, 34, 7, 9, 40, 11, 16, 34, 40, 13, 9, 12, 14])

ws = wb.create_sheet('not_in_corpus')
table(ws, list(sup[0].keys()), sup, [28, 34, 6, 50, 16, 26, 9, 60])

ws = wb.create_sheet('codebook')
CB = [
    ('study_id', '', 'Sequential identifier in manuscript order (S001-S126).'),
    ('bibkey', '', 'Citation key exactly as used in the manuscript and surveybib.bib.'),
    ('publication_type', 'Journal article / Conference paper / Book / Book chapter / Preprint / working paper', 'Taken from the BibTeX entry type and venue field, as used for the corpus-composition table.'),
    ('theme', 'Social media analysis / Financial market events / Cross-Domain Events and Causality', 'The three groups of the PRISMA flow diagram ("Included studies categorized by conceptual theme").'),
    ('prisma_box', '17 categories', 'The category box of the PRISMA diagram in which the study is counted. Each study is counted once, in the first section where it is analysed.'),
    ('xref_to', '', 'Further PRISMA category in which the study is also analysed. These are the "overlapping studies" shown in the diagram (2 in Financial signals, 1 in Cross-Domain Challenges, 5 in Causal Identification).'),
    ('section_no / section', '', 'Manuscript section that corresponds to the PRISMA category.'),
    ('method_class', 'C1-C6, F1-F5', 'Methodological class; assigned only where the study is listed in an appendix table.'),
    ('cited_in', '', 'All manuscript sections in which the study is cited.'),
    ('study_role', 'Core', 'Findings or contributions analysed to characterise the state of research.'),
    ('study_role', 'Foundational', 'Cited as the origin or standard reference of a definition, model or statistical method; included by the rule in Section 2.2 on foundational works.'),
    ('revision_status', 'restored to text in revision', 'Study belonged to the corpus counted in the PRISMA diagram; its citation had been lost from the text during revision and was restored.'),
    ('metadata_check', '', 'Result of the metadata check of 19 October 2026: OpenAlex work ID, and whether the DOI record was checked in Crossref. Bibliographic data are taken from surveybib_corrected.bib.'),
    ('note', '', 'Correction applied to the .bib entry, or a remaining point to verify.'),
    ('openalex_id', '', 'OpenAlex work identifier (https://openalex.org/<id>). Empty = not indexed in OpenAlex.'),
    ('in_search_window', 'yes / no', 'Publication year 2021-2026 (the database search window, Section 2.2). Only these studies enter the known-item retrieval check; older works are foundational works included through citation searching.'),
    ('validation_set', 'development / validation', 'Set of the known-item retrieval check (studies published 2021-2026 only; split at random within each PRISMA category, seed 20261019). Development titles were used to derive the query; validation studies were used only to evaluate it.'),
    ('retrieved_by_final_query', 'yes / no / not indexed in OpenAlex', 'Validation set only: whether the final query of the known-item retrieval check (sheet validation_result: term query with the date, document-type and subject-field restrictions) returned the study in OpenAlex on 20 October 2026. Empty for the development set, whose titles were used to build the query (development set: 42 of 44 retrieved).'),
]
table(ws, ['column', 'label', 'definition'], CB, [20, 60, 110])

# ---------------- class definitions (quoted from the revised manuscript, Sections 3.4 and 3.6)
_ms = open('ms_v2/manuscript_revised.tex', encoding='utf-8').read()
def _sentences(s, n):
    s = re.sub(r'\\textcolor\{blue\}\{', '', s); s = re.sub(r'\\citep\{[^}]*\}', '', s); s = re.sub(r'\s+', ' ', s).strip()
    return ' '.join(x.strip() for x in re.findall(r'.+?\.(?=\s|$)', s)[:n])
CD = []
for m in re.finditer(r'^\\textbf\{([CF][1-6]) - ([^}]*)\}(?:\\label\{[^}]*\})?\s*\n\s*\n(.+?)\n', _ms, re.M | re.S):
    code, name, para = m.groups()
    CD.append([code, name, _sentences(para, 2 if code.startswith('C') else 1),
               '3.4 Opinion Detection and Modeling' if code.startswith('C') else '3.6 Methodological Approaches in Financial Event Research',
               f'Table {code}', None])
assert [c[0] for c in CD] == ['C1', 'C2', 'C3', 'C4', 'C5', 'C6', 'F1', 'F2', 'F3', 'F4', 'F5'], [c[0] for c in CD]
ws = wb.create_sheet('class_definitions')
end = table(ws, ['class', 'name (manuscript heading)', 'definition (quoted from the manuscript)', 'manuscript section', 'appendix table', 'studies in register'],
            CD, [7, 44, 90, 40, 12, 12])
for i in range(2, end + 1):
    ws.cell(row=i, column=6, value=f'=COUNTIF(Data!$N$2:$N${DEND},A{i})').font = BF
ws.cell(row=end + 1, column=5, value='Total').font = BOLD
ws.cell(row=end + 1, column=6, value=f'=SUM(F2:F{end})').font = BOLD
notes_cd = [
    'Assignment: each of the 84 studies listed in the appendix tables is assigned to exactly one class (it appears in one appendix table). Studies were coded and classified by the first author according to their thematic focus and methodological characteristics, and the classifications were reviewed by the research supervisors; ambiguous cases were discussed jointly until agreement was reached (Section 2.2). The manuscript does not report an agreement statistic.',
    'The other 42 included studies (opinion and its manipulation, social-media and financial signals, cross-domain challenges, causal identification) are not assigned to a class (column method_class empty on Data).',
]
for i, s_ in enumerate(notes_cd):
    c = ws.cell(row=end + 3 + i, column=1, value=s_); c.font = Font(name=AR, size=9, italic=True)

ws = wb.create_sheet('Counts')
R = f'Data!$G$2:$G${DEND}'
ws.append(['Publication type', 'Reported (corpus-composition table)', 'Register', 'Difference'])
for c in range(1, 5):
    x = ws.cell(row=1, column=c); x.font, x.fill = HF, HFILL
for i, (t, v) in enumerate([('Journal article', 103), ('Conference paper', 7), ('Book', 3), ('Book chapter', 2), ('Preprint / working paper', 11)], 2):
    ws.cell(row=i, column=1, value=t).font = BF
    ws.cell(row=i, column=2, value=v).font = BLUE
    ws.cell(row=i, column=3, value=f'=COUNTIF({R},A{i})').font = BF
    ws.cell(row=i, column=4, value=f'=C{i}-B{i}').font = BF
ws.cell(row=7, column=1, value='Total').font = BOLD
for col in 'BCD':
    ws[f'{col}7'] = f'=SUM({col}2:{col}6)'; ws[f'{col}7'].font = BOLD
ws['A9'] = 'Theme'; ws['B9'] = 'Reported (PRISMA diagram)'; ws['C9'] = 'Register'; ws['D9'] = 'Difference'
for c in range(1, 5):
    x = ws.cell(row=9, column=c); x.font, x.fill = HF, HFILL
for i, (t, v) in enumerate([(T_SM, 59), (T_FIN, 40), (T_X, 27)], 10):
    ws.cell(row=i, column=1, value=t).font = BF
    ws.cell(row=i, column=2, value=v).font = BLUE
    ws.cell(row=i, column=3, value=f'=COUNTIF(Data!$I$2:$I${DEND},A{i})').font = BF
    ws.cell(row=i, column=4, value=f'=C{i}-B{i}').font = BF
ws['A13'] = 'Total'; ws['A13'].font = BOLD
for col in 'BCD':
    ws[f'{col}13'] = f'=SUM({col}10:{col}12)'; ws[f'{col}13'].font = BOLD
ws['A15'] = 'Study role'; ws['C15'] = 'Register'
for c in (1, 3):
    x = ws.cell(row=15, column=c); x.font, x.fill = HF, HFILL
for i, t in enumerate(['Core', 'Foundational'], 16):
    ws.cell(row=i, column=1, value=t).font = BF
    ws.cell(row=i, column=3, value=f'=COUNTIF(Data!$Q$2:$Q${DEND},A{i})').font = BF
ws['A19'] = 'Publication years (min / max)'; ws['A19'].font = BF
ws['B19'] = 'Reported: 1969-2026'; ws['B19'].font = BLUE
ws['C19'] = f'=MIN(Data!$D$2:$D${DEND})&"-"&MAX(Data!$D$2:$D${DEND})'; ws['C19'].font = BF
ws['A21'] = 'OpenAlex (metadata check and retrieval check)'; ws['C21'] = 'Register'
for c in (1, 3):
    x = ws.cell(row=21, column=c); x.font, x.fill = HF, HFILL
VW, VU, VR = f'Data!$V$2:$V${DEND}', f'Data!$U$2:$U${DEND}', f'Data!$W$2:$W${DEND}'
for i, (lab, f) in enumerate([
        ('Included studies published 2021-2026 (search window)', f'=COUNTIF({VW},"yes")'),
        ('  of which indexed in OpenAlex', f'=COUNTIFS({VW},"yes",{VU},"<>")'),
        ('Studies (of 126) indexed in OpenAlex (metadata check)', f'=COUNTIF({VU},"<>")')], 22):
    ws.cell(row=i, column=1, value=lab).font = BF
    c = ws.cell(row=i, column=3, value=f); c.font = BF
    if 'Recall' in lab:
        c.number_format = '0.0%'
ws['A30'] = 'Blue = value reported in the manuscript; black = computed from the Data sheet.'; ws['A30'].font = Font(name=AR, size=9, italic=True)
for c, w in zip('ABCD', [40, 34, 12, 12]):
    ws.column_dimensions[c].width = w

ws = wb.create_sheet('PRISMA_check')
hd = ['Theme', 'PRISMA category', 'Diagram: n', 'Diagram: overlapping', 'Register: counted here', 'Register: overlapping', 'Register: n', 'Match']
for c, h in enumerate(hd, 1):
    x = ws.cell(row=1, column=c, value=h); x.font, x.fill = HF, HFILL; x.alignment = WR
for i, (box, (theme, *_)) in enumerate(BOXES.items(), 2):
    n, ov = DIAGRAM[box]
    ws.cell(row=i, column=1, value=theme).font = BF
    ws.cell(row=i, column=2, value=box).font = BF
    ws.cell(row=i, column=3, value=n).font = BLUE
    ws.cell(row=i, column=4, value=ov).font = BLUE
    ws.cell(row=i, column=5, value=f'=COUNTIF(Data!$J$2:$J${DEND},B{i})').font = BF
    ws.cell(row=i, column=6, value=f'=SUMPRODUCT(--ISNUMBER(SEARCH(B{i},Data!$K$2:$K${DEND})))').font = BF
    ws.cell(row=i, column=7, value=f'=E{i}+F{i}').font = BF
    ws.cell(row=i, column=8, value=f'=IF(AND(G{i}=C{i},F{i}=D{i}),"yes","NO")').font = BF
last = 1 + len(BOXES)
ws.cell(row=last + 1, column=2, value='Total studies (each counted once)').font = BOLD
ws.cell(row=last + 1, column=3, value=126).font = BLUE
ws.cell(row=last + 1, column=5, value=f'=SUM(E2:E{last})').font = BOLD
ws.cell(row=last + 1, column=8, value=f'=IF(E{last + 1}=C{last + 1},"yes","NO")').font = BOLD
for c, w in zip('ABCDEFGH', [30, 50, 11, 13, 14, 14, 11, 8]):
    ws.column_dimensions[c].width = w

ws = wb.create_sheet('search_concepts_table2')
SB = [
    ('Social Media Manipulation', 'misinformation, disinformation, fake news, propaganda, opinion manipulation, influence operations',
     '("social media" OR "online social networks") AND (misinformation OR disinformation OR "fake news" OR propaganda OR "opinion manipulation")', 1667),
    ('Opinion Dynamics in Networks', 'opinion dynamics, social influence, information diffusion, consensus formation, bounded confidence',
     '("opinion dynamics" OR "information diffusion" OR "social influence") AND ("social networks" OR "online social networks")', 762),
    ('Event Detection in Social Media', 'event detection, burst detection, topic detection, social media events, event extraction',
     '("event detection" OR "bursty events" OR "event extraction") AND ("social media" OR "online social networks")', 172),
    ('Fake News and Propaganda Detection', 'fake news detection, propaganda detection, rumor detection, misinformation classification',
     '("fake news detection" OR "propaganda detection" OR "rumor detection") AND ("social media" OR "online platforms")', 382),
    ('Financial Events and Market Responses', 'event studies, abnormal returns, volatility, structural changes, high-frequency market responses',
     '("event study" OR "financial event" OR "market event" OR "information event") AND ("abnormal return" OR volatility OR "structural break" OR "market reaction")', 1066),
    ('Social Media and Financial Markets', 'investor sentiment, investor attention, misinformation, market manipulation, pump and dump, market reactions',
     '("social media" OR Twitter OR Reddit OR Telegram) AND ("stock market" OR "financial market" OR cryptocurrency) AND ("market reaction" OR "market impact" OR "price impact" OR "return predictability" OR "market manipulation")', 530),
]
end = table(ws, ['concept', 'keywords and synonyms', 'representative Boolean string (as reported)', 'records (as reported)'], SB, [30, 50, 90, 12])
ws.cell(row=end + 1, column=3, value='Total').font = BOLD
ws.cell(row=end + 1, column=4, value=f'=SUM(D2:D{end})').font = BOLD
ws.cell(row=end + 3, column=1, value='Source: search-concepts table (tab:search_concepts), Section 2.2. They are not the query of the known-item retrieval check.').font = Font(name=AR, size=9, italic=True)

# ---------------- known-item retrieval check: one protocol (val/PROTOCOL_final.md)
ws = wb.create_sheet('search_protocol')
for i, line in enumerate(open('val/PROTOCOL_final.md', encoding='utf-8').read().splitlines(), 1):
    c = ws.cell(row=i, column=1, value=line or None)
    c.font = BOLD if line.startswith('#') else BF
    c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 130

def pct(d):
    return f"{d['retrieved']}/{d['n']} = {d['share']:.1%} (95% CI {d['ci95'][0]:.1%}-{d['ci95'][1]:.1%})"
VV = V2['variants']
FE = FINAL
FIELD_OF = {r['bibkey']: r['field'] for r in FIELDS}
ws = wb.create_sheet('validation_result')
fin_missed = FINAL['validation_missed']
V1R = VAL_LOG['results']
info = [
    ('protocol', 'search_protocol (PROTOCOL.md). Term query, threshold, date and document-type restrictions fixed on 19 October 2026 before any count; subject-field restriction added on 20 October 2026 after earlier runs; a language filter was tested and dropped (sheet protocol_history).'),
    ('what the check measures', 'Known-item retrieval within the previously selected corpus. Because both the development and validation samples originate from that corpus, the results are conditional on the original study selection and cannot establish the completeness, representativeness or sensitivity of the original systematic search.'),
    ('database', 'OpenAlex, title search (title.search). Term query frozen on 19 October 2026; final run on 20 October 2026 (run_final.py).'),
    ('date restriction', VAL_LOG['date_filter'] + ' - one year before the 2021-2026 window of the included studies, because OpenAlex dates some works by their first online or preprint version (among the 82 studies, arnold2022attention, published in 2022, is dated 2020 in OpenAlex; no validation study is dated before 2021).'),
    ('document-type restriction', VAL_LOG['type_filter'] + ' - all other types (journal articles, conference papers, books, book chapters, preprints) are kept.'),
    ('subject-field restriction', 'Not an eligibility criterion; used to reduce the number of records. OpenAlex field of the primary topic, kept if it is the field of at least one development study: ' + '; '.join(f"{v[0]} ({v[1]} development studies)" for v in FE['fields_kept'].values()) + '. Chosen after the fields of the validation studies had been inspected, so not independently validated.'),
    ('language restriction', 'none. OpenAlex leaves some works without a language, and a language filter would exclude two included studies (cui2025toward, goldsmith2025causal); the English-language criterion is applied at screening.'),
    ('N (works with the date and type restrictions; used for the ratios)', f"{VAL_LOG['N_works']:,}"),
    ('threshold R (chosen on the development set)', VAL_LOG['R']),
    ('frozen term query (27 terms)', VAL_LOG['query']),
    ('term query SHA-256', VAL_LOG['query_sha256'] + ' (documents the exact query run; it does not by itself show that the configuration was chosen independently of the validation results)'),
    ('records returned', f"{FE['records']:,} (without the subject-field restriction on the same day: {FE['records_term_query']:,}; removed by the restriction: {FE['records_term_query'] - FE['records']:,}, {(FE['records_term_query'] - FE['records']) / FE['records_term_query']:.1%})"),
    ('records not tagged as English (descriptive)', f"{FE['records_not_tagged_english']:,} ({FE['records_not_tagged_english'] / FE['records']:.1%}; another language or no language assigned)"),
    ('validation set, all (primary outcome)', pct(FE['validation_all'])),
    ('validation set, indexed in OpenAlex', pct(FE['validation_indexed'])),
    ('development set', pct(FE['development'])),
    ('validation studies indexed in OpenAlex but not retrieved (7)', '; '.join(FINAL['validation_missed'])),
    ('validation studies not indexed in OpenAlex (2)', '; '.join(FINAL['validation_not_indexed'])),
    ('interpretation', 'The subject-field restriction did not remove any of the known included studies retrieved by the original query: the original pre-specified configuration retrieved the same validation studies (sheet protocol_history). The check does not show that the records removed by the restriction contain no other relevant studies, and it cannot be used as evidence of the absence of selection bias.'),
    ('method note', 'Offline matching of the development titles was used exclusively for threshold selection; all retrieval outcomes were determined from the actual OpenAlex search results.'),
    ('field check', 'All 36 indexed validation studies have their primary topic in one of the fields kept (sheet openalex_fields).'),
]
end = table(ws, ['item', 'value'], [list(x) for x in info], [42, 140])
r0 = end + 2
ws.cell(row=r0, column=1, value='Threshold selection on the development set (offline title match)').font = BOLD
table(ws, ['R', 'terms in query', 'development titles matched'], [[d['R'], d['n_terms'], d['dev_match_offline']] for d in VAL_LOG['threshold_trace']], [42, 14, 24], start=r0 + 1)
r1 = r0 + 3 + len(VAL_LOG['threshold_trace'])
ws.cell(row=r1, column=1, value='Validation set: per-study result').font = BOLD
vrows = [[k, r['prisma_box'], r['openalex_id'] or None, FIELD_OF.get(k), val_final(k)] for k, r in VAL_RET.items() if r['set'] == 'validation']
assert sum(x[4] == 'yes' for x in vrows) == FE['validation_all']['retrieved']
table(ws, ['bibkey', 'PRISMA category', 'OpenAlex ID', 'OpenAlex field', 'retrieved'], vrows, [42, 50, 14, 32, 16], start=r1 + 1)
r4 = r1 + 3 + len(vrows)
ws.cell(row=r4, column=1, value='Additional descriptive count (not part of the protocol; nothing filtered): records of the final query by OpenAlex source type and work type, 20 October 2026 (count_by_source_type.py)').font = BOLD
TS = TYPES_SUM
s_end = table(ws, ['source type (primary location)', 'records', 'share'], [[x['type'], x['records'], x['records'] / TS['records']] for x in TS['by_source_type']], [42, 50, 14], start=r4 + 1)
w_end = table(ws, ['work type', 'records', 'share'], [[x['type'], x['records'], x['records'] / TS['records']] for x in TS['by_work_type']], [42, 50, 14], start=s_end + 2)
for rr in range(r4 + 2, w_end + 1):
    ws.cell(row=rr, column=2).number_format = '#,##0'
    ws.cell(row=rr, column=3).number_format = '0.0%'
jr = next(x['records'] for x in TS['by_source_type'] if x['type'] == 'journal')
inel = sum(x['records'] for x in TS['by_work_type'] if x['type'] in ('libguides', 'software', 'book-review', 'reference-entry', 'supplementary-materials', 'standard', 'software-paper', 'conference-abstract'))
tnotes = [
    f"Included studies 2021-2026 by OpenAlex source type: development set journal 38, repository 2, conference 1, ebook platform 1, book series 1, no source 1; validation set journal 32, repository 2, no source 2, not indexed 2.",
    f"A journal-only restriction (not applied) would keep {jr:,} records ({jr / TS['records']:.1%}) but lose the four retrieved validation studies without a journal source (" + ', '.join(f'{k}: {v}' for k, v in TS['validation_retrieved_not_in_journal'].items()) + "): validation 25/38 (65.8%). It would also exclude conference papers and preprints, which the review admits.",
    f"Work types that the review would not include and that the document-type restriction did not exclude (libguides, software, book reviews, reference entries, supplementary materials, standards, software papers, conference abstracts) account for {inel:,} records ({inel / TS['records']:.1%}).",
]
for i, s in enumerate(tnotes):
    c = ws.cell(row=w_end + 2 + i, column=1, value=s); c.font = Font(name=AR, size=9, italic=True)
ws.freeze_panes = None

ws = wb.create_sheet('validation_terms')
VT = [OrderedDict([('term', r['term']), ('kind', r['kind']), ('source', r['source']), ('development titles (of 44)', int(r['df_dev'])),
                   ('development categories', r['dev_categories']), ('OpenAlex title count', int(r['oa_count'])),
                   ('over-representation ratio', float(r['ratio'])), ('in frozen query', 'yes' if r['in_query'] == 'True' else 'no')])
      for r in VAL_TERMS]
end = table(ws, list(VT[0].keys()), VT, [26, 8, 20, 12, 50, 14, 14, 10])
ws.cell(row=end + 2, column=1, value='Over-representation ratio = (development titles containing the term / 44) / (OpenAlex title count / N), with the date and document-type restrictions. It measures enrichment of the term in the development set relative to OpenAlex; it is used for query construction and is not an estimate of search precision or specificity. Terms with ratio >= R = 100 were kept; a phrase was dropped when one of its words was itself kept (e.g. "echo chamber" is covered by "echo" and "chamber").').font = Font(name=AR, size=9, italic=True)

ws = wb.create_sheet('protocol_history')
FROW = dict(FE, missed_beyond_v1=[])
HIST = [('term query without subject-field restriction (rules fixed before any count)', '2026-10-19', 'term query; dates 2020-2026; ineligible types excluded', 'original pre-specified configuration (audit trail)', VV['v1']),
        ('+ language filter', '2026-10-20', 'as above + language:en', 'dropped (lost two studies without a language)', VV['A']),
        ('+ language and subject-field filters', '2026-10-20', 'as above + fields of the development studies', 'dropped (language filter)', VV['E']),
        ('final: + subject-field restriction, no language filter', '2026-10-20', 'term query; dates 2020-2026; ineligible types excluded; fields of the development studies', 'final configuration (reported filtered retrieval result)', FROW),
        ('phrases instead of generic terms, K = 50,000', '2026-10-20', 'language:en; every term with > 50,000 title matches replaced by its candidate phrases', 'rejected (lost validation studies)', VV['B']),
        ('phrases instead of generic terms, K = 10,000', '2026-10-20', 'language:en; every term with > 10,000 title matches replaced by its candidate phrases', 'rejected (lost validation studies)', VV['D'])]
def lost_txt(v):
    if v['missed_beyond_v1'] is None:
        return f"{VV['v1']['validation_all']['retrieved'] - v['validation_all']['retrieved']} studies"
    return '; '.join(v['missed_beyond_v1']) or '-'
hrows = [[a, d, f, s, v['records'], pct(v['development']), pct(v['validation_all']), pct(v['validation_indexed']), lost_txt(v)] for a, d, f, s, v in HIST]
end = table(ws, ['query', 'date', 'restrictions / change', 'status', 'records', 'development set (44)', 'validation set, all (38)', 'validation set, indexed (36)',
                 'validation studies lost compared with the first configuration'], hrows, [34, 11, 46, 22, 10, 26, 26, 26, 40])
for rr in range(2, 2 + len(hrows)):
    ws.cell(row=rr, column=5).number_format = '#,##0'
hnotes = [
    'The term query (27 terms, R = 100) is the same in every row; only the filters change, except in the two rejected rows. Each query was run once. Dated protocol documents with their SHA-256: history/ in the repository (PROTOCOL_v1.md, PROTOCOL_v2.md, PROTOCOL_final_prerun.md and the manifests).',
    'The subject-field restriction was added after the first run, so it was not specified in advance. The field rule (fields of the development studies only) was written before any field data were retrieved; the final configuration was hashed before its run but after the fields of the validation studies were known (show_fields.py), so the restriction is not independently validated.',
    'Language check (20 October 2026): OpenAlex assigns no language to cui2025toward (W4415736465) and goldsmith2025causal (W7106255595), the two validation studies lost with the language filter. The language filter was therefore dropped; the final query retrieves the same validation studies as the first configuration.',
    f"Records on 20 October 2026: without the subject-field restriction the query returned {FE['records_term_query']:,} records (19 October 2026: 528,412; OpenAlex is updated continuously). The reduction by the restriction is computed between the two counts of 20 October 2026.",
    'Records of the first configuration by language (20 October 2026): tagged as English 427,734; not tagged as English (another language or no language assigned) 100,678 (19.1%). Dated 2020: 64,716 (12.2%); with a start date of 2021-01-01 the first configuration would return 463,696 records, with the same validation result.',
    'Confidence intervals are Wilson score intervals.',
]
for i, s in enumerate(hnotes):
    c = ws.cell(row=end + 2 + i, column=1, value=s); c.font = Font(name=AR, size=9, italic=True)
r3 = end + 3 + len(hnotes)
ws.cell(row=r3, column=1, value='Records of the query without the subject-field restriction, by OpenAlex field of the primary topic (20 October 2026)').font = BOLD
BYF = FE['records_term_query_by_field']
fend = table(ws, ['OpenAlex field', 'records', 'kept by the field restriction'], [[x['field'], x['records'], 'yes' if x['kept'] else 'no'] for x in BYF], [34, 11, 46], start=r3 + 1)
ws.cell(row=fend + 1, column=1, value='total').font = BOLD
ws.cell(row=fend + 1, column=2, value=f'=SUM(B{r3 + 2}:B{fend})')
ws.cell(row=fend + 2, column=1, value='kept by the field restriction').font = BOLD
ws.cell(row=fend + 2, column=2, value=f'=SUMIF(C{r3 + 2}:C{fend},"yes",B{r3 + 2}:B{fend})')
for rr in range(r3 + 2, fend + 3):
    ws.cell(row=rr, column=2).number_format = '#,##0'
ws.freeze_panes = None

ws = wb.create_sheet('openalex_fields')
fh = ['bibkey', 'set', 'prisma_box', 'openalex_id', 'field', 'subfield', 'topic']
fend = table(ws, ['bibkey', 'set', 'PRISMA category', 'OpenAlex ID', 'field (primary topic)', 'subfield', 'primary topic'],
             [[r[h] or None for h in fh] for r in FIELDS], [28, 12, 46, 13, 32, 40, 44])
r0 = fend + 2
ws.cell(row=r0, column=1, value='Studies per field (formulas)').font = BOLD
flist = ['Social Sciences', 'Economics, Econometrics and Finance', 'Computer Science', 'Physics and Astronomy', 'Decision Sciences', 'not indexed']
cnt_rows = [[f, f'=COUNTIFS($E$2:$E${fend},A{r0 + 2 + i},$B$2:$B${fend},"development")', f'=COUNTIFS($E$2:$E${fend},A{r0 + 2 + i},$B$2:$B${fend},"validation")']
            for i, f in enumerate(flist)]
cend = table(ws, ['field', 'development', 'validation'], cnt_rows, [28, 12, 46], start=r0 + 1)
ws.cell(row=cend + 1, column=1, value='total').font = BOLD
ws.cell(row=cend + 1, column=2, value=f'=SUM(B{r0 + 2}:B{cend})')
ws.cell(row=cend + 1, column=3, value=f'=SUM(C{r0 + 2}:C{cend})')
ws.cell(row=cend + 3, column=1, value='OpenAlex field, subfield and topic of each study\'s primary topic, read with show_fields.py on 20 October 2026 (assigned automatically by OpenAlex). '
        'The subject-field restriction keeps the fields of the development studies; every indexed validation study falls in one of them.').font = Font(name=AR, size=9, italic=True)
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:G{fend}'

# ---------------- metadata corrections
BC = list(csv.DictReader(open('bib_corrections_log.csv', encoding='utf-8')))
ws = wb.create_sheet('bib_corrections')
table(ws, ['bibkey', 'field', 'old', 'new', 'reason', 'source'], BC, [26, 12, 50, 50, 50, 40])

ws = wb.create_sheet('decisions_log')
DL = [
    ('2026-10-18', 'Register compiled from the final manuscript and surveybib.bib', 'Each included work linked to its PRISMA category, section and methodological class.'),
    ('2026-10-18', 'Corpus membership defined by the PRISMA category boxes', 'A cited work belongs to the corpus if it is analysed in one of the sections counted in the PRISMA diagram (3.2-3.4, 3.6, 4.1.1, 4.1.2, 4.2, 4.2.1). Works cited only in Sections 1-2 (19) are listed on not_in_corpus.'),
    ('2026-10-18', 'de2026transformer removed from Class 2 (text and Table C2)', 'Not counted in the PRISMA diagram (Class 2, n = 8); it reviews text generation rather than manipulation detection.'),
    ('2026-10-18', 'tong2022assessing, seetharam2024headlines, ghazouani2025virtual restored to Section 4.2.1', 'Counted in the PRISMA diagram (Causal Identification, n = 19) but their citations had been lost from the text during revision. Confirmed by the author; descriptions checked against full texts (seetharam, ghazouani) and abstract (tong).'),
    ('2026-10-18', 'wang2022causal restored to Section 4.2', 'Counted in the PRISMA diagram (Cross-Domain Challenges, n = 6) but its citation had been lost from the text. Confirmed by the author; description checked against the full text.'),
    ('2026-10-18', 'Overlapping studies counted once', 'Each study is counted in the first PRISMA category where it is analysed; later appearances are recorded in xref_to, reproducing the overlaps shown in the diagram.'),
    ('2026-10-18', 'CORPUS: 126 studies; all PRISMA category counts and the publication-type composition reproduced', 'See PRISMA_check and Counts.'),
    ('2026-10-19', 'Metadata of all 126 studies checked against OpenAlex, and their DOIs against Crossref', 'verify_corpus.py; 122 of 126 studies indexed in OpenAlex (not indexed: musch2003psychology, cayuela2025influence, kumarsocial; for campbell1997econometrics only a journal item about the book).'),
    ('2026-10-19', '17 corrections applied to surveybib.bib (surveybib_corrected.bib), including 10 DOIs', 'Five DOIs resolved to other articles, four were not registered in Crossref, one was truncated; see bib_corrections.'),
    ('2026-10-19', 'Validated verification search: protocol fixed before any OpenAlex count (val/PROTOCOL.md, MANIFEST.txt)',
     'Objective approach (Hausner et al. 2012): stratified random split (seed 20261019) into development (44) and validation (38) sets; terms from development titles (>= 2 titles) and Table 2 keywords; specificity by over-representation relative to OpenAlex; threshold chosen on the development set; one frozen title query run once.'),
    ('2026-10-19', 'Validated search, result: R = 100, 27 terms, 528,412 records; validation 29/38 (76.3%, 95% CI 60.8-87.0%), 29/36 indexed (80.6%, 95% CI 65.0-90.2%); development 42/44 (95.5%)',
     'Reported as obtained. The query is broad because the corpus spans three domains; the validation studies not retrieved use none of the core terms in their titles (fact-checkers, checkworthiness, emotion, tweets, NFT promotion scams) or, for bradley2024place, were not returned by OpenAlex despite "Reddit" in the title.'),
    ('2026-10-19', 'Deviation from the protocol (technical): matching step re-executed (run_step3.py)',
     'In the first execution the register file used had no OpenAlex identifiers, so no included study could be matched (0 of 82, all reported as not indexed). The matching was re-executed with the same frozen query (SHA-256 checked) and the identifiers of this register; terms, threshold and query unchanged.'),
    ('2026-10-19', 'After review: the OpenAlex check is reported as a supplementary, retrospective retrievability check',
     'It tests whether the included studies can be retrieved by a query not tailored to them; it does not reproduce the original search and does not estimate its sensitivity. The query is broad (528,412 records).'),
    ('2026-10-19', 'Register synchronised with the revised manuscript (citation inventory regenerated; cited_in columns updated)',
     'The revised manuscript cites 146 keys: the 126 corpus studies (including the four restored citations) and the 20 works on not_in_corpus; de2026transformer is no longer cited. All keys are in surveybib_corrected.bib. The OpenAlex check is named a retrospective known-item retrieval check.'),
    ('2026-10-19', 'Sheet class_definitions added',
     'Definitions of the methodological classes C1-C6 and F1-F5 quoted from the manuscript (Sections 3.4 and 3.6), with the assignment procedure of Section 2.2 and the number of studies per class (84 in total).'),
    ('2026-10-19', 'hausner2012routine added to the bibliographic database and to not_in_corpus',
     'Method reference cited in the new subsection Reproducibility of the Review; logged as a new entry on bib_corrections. not_in_corpus now lists 20 works.'),
    ('2026-10-19', 'Table 2 of the manuscript (search concepts and record counts of the primary search) left unchanged',
     'Table 2 documents the search that produced the 4,579 records of the PRISMA diagram; the validated verification search was run later, in OpenAlex, and is reported separately (make_validation_table.py).'),
    ('2026-10-20', 'Records of the version-1 query counted by language and by year (descriptive)',
     'Tagged as English 427,734 and not tagged as English 100,678 of 528,412; dated 2020: 64,716 (sheet validation_result). The run itself is unchanged.'),
    ('2026-10-20', 'Protocol version 2 written after the version-1 result (PROTOCOL_v2.md, MANIFEST_v2.txt): language:en and a specificity cap K; variants B (K = 50,000) and D (K = 10,000) frozen and run once',
     'B: 297,641 records, validation 24/38 (63.2%, 95% CI 47.3-76.6%); D: 49,580 records, validation 12/38 (31.6%, 95% CI 19.1-47.5%). Narrowing the query reduced the records but also the retrieved studies.'),
    ('2026-10-20', 'Variant A added by the author after B and D were run (addendum of PROTOCOL_v2.md): version-1 query unchanged plus language:en only',
     'Run once (run_v2A.py; the version-1 query is checked against its SHA-256): 427,734 records; development 42/44; validation 27/38 (71.1%, 95% CI 55.2-83.0%), 27/36 indexed (75.0%, 95% CI 58.9-86.2%). Lost compared with version 1: cui2025toward and goldsmith2025causal, to which OpenAlex assigns no language.'),
    ('2026-10-20', 'Reporting in the manuscript (Reproducibility of the Review)',
     'Version 1 is reported as the pre-specified result; the English restriction (variant A) and the phrase-based variants (B, D) are reported as later restrictions, with their records and retrieved studies.'),
    ('2026-10-20', 'Variant E added by the author after A was run (addendum 2 of PROTOCOL_v2.md): variant A plus a filter on the OpenAlex fields of the primary topic, keeping every field of at least one development study',
     'The version-1 query searches all subject areas (e.g. chemistry, medicine). Fields kept: Social Sciences, Economics/Econometrics/Finance, Computer Science, Physics and Astronomy, Decision Sciences. Run once (run_v2E.py): 247,876 records (variant A: 427,734); development 42/44; validation 27/38 (71.1%, 95% CI 55.2-83.0%), the same studies as variant A. The field filter removed 179,858 records (42.0%) and none of the included studies retrieved by variant A.'),
    ('2026-10-20', 'Reporting in the manuscript updated',
     'Reproducibility of the Review reports version 1 (pre-specified), the English restriction and the field restriction (records 427,734 and 247,876, both 27/38), and the phrase-based variants B and D (24/38, 12/38).'),
    ('2026-10-20', 'Primary topics of the 82 studies read from OpenAlex (show_fields.py; sheet openalex_fields)',
     'Development set: Social Sciences 23, Economics/Econometrics/Finance 15, Computer Science 3, Physics and Astronomy 2 (ge2024novel: Complex Network Analysis Techniques; zhang4450056integrated: Opinion Dynamics and Social Influence), Decision Sciences 1. '
     'Validation set: Social Sciences 21, Economics/Econometrics/Finance 9, Computer Science 4, Physics and Astronomy 1, Decision Sciences 1, not indexed 2. All indexed validation studies fall in the five fields kept; no difference from the fields recorded in the variant-E run.'),
    ('2026-10-20', 'One protocol (PROTOCOL.md; sheet search_protocol) replaces the dated protocol documents as the description of the check',
     'It combines the rules fixed on 19 October 2026 with the language and field filters added on 20 October 2026 (variant E), and records when each rule was fixed. The reported result is that of the final query (variant E): 247,876 records; validation 27/38. The first run and the rejected variants are on protocol_history; the dated documents remain in the repository (history/).'),
    ('2026-10-20', 'Language filter dropped; final query = term query with the date, document-type and subject-field filters (PROTOCOL.md, hashed in MANIFEST_final.txt before the run)',
     'The language filter removes works without an assigned language, including two included studies. Run once (run_final.py): 311,556 records (term query on the same day: 528,902); not tagged as English 63,578; development 42/44; validation 29/38 (76.3%, 95% CI 60.8-87.0%), the same studies as the first run. Supersedes the previous entry.'),
    ('2026-10-20', 'Reporting revised after methodological review: reference query = reported result; field restriction = post-hoc sensitivity analysis',
     'The field-restricted query was specified after the validation results and the fields of the validation studies were known, so it is not an independent validation. The reference query (rules fixed before execution; 29/38; 528,412 records) remains the independently evaluated result. PROTOCOL.md revised editorially (known-item statement; "Retrieval restrictions"; "Term enrichment and threshold selection"; SHA-256 note; one run per configuration; 7 indexed-not-retrieved vs 2 not indexed). No rule, query or result changed. Data column renamed retrieved_by_reference_query (all 82 studies).'),
    ('2026-10-20', "Author's decision: the latest run is the reported result",
     'Reported result: final query (term query with date, document-type and subject-field restrictions; no language restriction): 311,556 records; validation 29/38. The manuscript, the table and this register state that the subject-field restriction was chosen after the fields of the validation studies were inspected and is not independently validated. The final configuration is used for the reported filtered retrieval result, while the original pre-specified configuration and all intermediate runs are retained as part of the methodological audit trail. They are documented on protocol_history and in the repository (history/). Replaces the reporting described in the previous entry. Data column: retrieved_by_final_query (validation set).'),
    ('2026-10-20', 'Additional descriptive count: records of the final query by OpenAlex source type and work type (count_by_source_type.py); nothing filtered, protocol and result unchanged',
     'Journal 175,350 (56.3%), repository 54,218, unknown 27,451, conference 25,301, book series 14,660, ebook platform 14,576. A journal-only restriction was considered and not applied: it would lose four retrieved validation studies (cui2025toward, huang2025hybrid without a source; fan2024measuring, goldsmith2025causal in a repository; validation 25/38) and contradict the eligibility criteria (conference papers and preprints are admitted). Ineligible work types not excluded by the document-type restriction account for 5,342 records (1.7%).'),
]
table(ws, ['date', 'decision', 'rationale'], DL, [12, 60, 110])

for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if c.font is None or c.font.name != AR:
                c.font = Font(name=AR, size=c.font.size if c.font else 9, bold=c.font.bold if c.font else False,
                              color=c.font.color if c.font else None, italic=c.font.italic if c.font else False)
wb.save(OUT)
print('saved', OUT, len(rows), len(sup))

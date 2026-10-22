#!/usr/bin/env python3
"""Offline verification of this repository (no network, about one minute).

  pip install -r requirements.txt
  python3 verify_repository.py

Checks
  1. SHA-256 hashes written before each run (retrieval_check/history/MANIFEST*.txt) against the files.
  2. The development/validation split is reproduced exactly from Supplementary Table S1 (split_corpus.py).
  3. The corrected bibliography and its log are reproduced exactly from the original .bib (fix_bib.py).
  4. The LaTeX table of the retrieval check is reproduced exactly from the run outputs (make_validation_table.py).
  5. Wilson 95% confidence intervals are recomputed from the counts.
  6. Supplementary Table S1: 126 studies, PRISMA category counts, validation set = split.csv, per-study
     retrieval = run outputs.
  7. The result in PROTOCOL.md (Section 9) and in the summaries equals the run outputs.
  8. The per-study outputs of the final run agree with the summary.
Each check prints OK or FAIL; the last line is ALL CHECKS PASSED or SOME CHECKS FAILED.
"""
import csv, hashlib, json, math, os, re, shutil, subprocess, sys, tempfile

ROOT = os.path.dirname(os.path.abspath(__file__))
RC = os.path.join(ROOT, 'retrieval_check')
HIST = os.path.join(RC, 'history')
S1 = os.path.join(ROOT, 'Supplementary_Table_S1_Corpus_Register.xlsx')
results = []


def check(cond, msg):
    results.append(bool(cond))
    print(('  OK    ' if cond else '  FAIL  ') + msg)


def warn(msg):
    print('  NOTE  ' + msg)


def sha(path):
    return hashlib.sha256(open(path, 'rb').read()).hexdigest()


def same(a, b):
    return os.path.exists(a) and os.path.exists(b) and open(a, 'rb').read() == open(b, 'rb').read()


def run(cmd, cwd):
    p = subprocess.run([sys.executable] + cmd, cwd=cwd, capture_output=True, text=True)
    if p.returncode:
        print(p.stdout[-800:], p.stderr[-800:])
    return p.returncode == 0


def wilson(k, n, z=1.96):
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return p, c - h, c + h


try:
    import openpyxl
except ImportError:
    sys.exit('openpyxl is missing: pip install -r requirements.txt')

# 1. hashes ---------------------------------------------------------------------------------------
print('1. SHA-256 hashes written before each run')
RENAMED = {('MANIFEST.txt', 'PROTOCOL.md'): os.path.join(HIST, 'PROTOCOL_v1.md'),
           ('MANIFEST_v2.txt', 'PROTOCOL_v2.md'): os.path.join(HIST, 'PROTOCOL_v2.md'),
           ('MANIFEST_final.txt', 'PROTOCOL.md'): os.path.join(HIST, 'PROTOCOL_final_prerun.md')}
for man in ('MANIFEST.txt', 'MANIFEST_v2.txt', 'MANIFEST_final.txt'):
    last = {}
    for line in open(os.path.join(HIST, man), encoding='utf-8'):
        m = re.match(r'^([0-9a-f]{64})  (\S+)$', line.strip())
        if m:
            last[m.group(2)] = m.group(1)          # a file listed twice: the later hash is the current one
    for name, h in last.items():
        path = RENAMED.get((man, name), os.path.join(RC, name))
        check(os.path.exists(path) and sha(path) == h, f'{man}: {os.path.relpath(path, ROOT)}')

tmp = tempfile.mkdtemp(prefix='verify_')
try:
    # 2. split ------------------------------------------------------------------------------------
    print('2. Development/validation split')
    d = os.path.join(tmp, 'split'); os.makedirs(d)
    shutil.copy(os.path.join(RC, 'split_corpus.py'), d)
    ok = run(['split_corpus.py', S1], d)
    check(ok and same(os.path.join(d, 'split.csv'), os.path.join(RC, 'split.csv')), 'split_corpus.py reproduces retrieval_check/split.csv exactly')

    # 3. bibliography -----------------------------------------------------------------------------
    print('3. Bibliographic corrections')
    d = os.path.join(tmp, 'bib'); os.makedirs(d)
    shutil.copy(os.path.join(ROOT, 'bibliography', 'surveybib.bib'), d)
    ok = run([os.path.join(ROOT, 'scripts', 'fix_bib.py')], d)
    check(ok and same(os.path.join(d, 'surveybib_corrected.bib'), os.path.join(ROOT, 'bibliography', 'surveybib_corrected.bib')),
          'fix_bib.py reproduces bibliography/surveybib_corrected.bib exactly')
    check(ok and same(os.path.join(d, 'bib_corrections_log.csv'), os.path.join(ROOT, 'bibliography', 'bib_corrections_log.csv')),
          'fix_bib.py reproduces bibliography/bib_corrections_log.csv exactly')

    # 4. LaTeX table ------------------------------------------------------------------------------
    print('4. LaTeX table of the retrieval check')
    d = os.path.join(tmp, 'table'); os.makedirs(d)
    for sub in ('run_output', 'run_output_final', 'run_output_v2'):
        shutil.copytree(os.path.join(RC, sub), os.path.join(d, sub))
    ok = run([os.path.join(ROOT, 'scripts', 'make_validation_table.py'), 'run_output'], d)
    check(ok and same(os.path.join(d, 'validation_table.tex'), os.path.join(RC, 'validation_table.tex')),
          'make_validation_table.py reproduces retrieval_check/validation_table.tex exactly')
finally:
    shutil.rmtree(tmp, ignore_errors=True)

# 5. confidence intervals ---------------------------------------------------------------------------
print('5. Wilson 95% confidence intervals')
F = json.load(open(os.path.join(RC, 'run_output_final', 'results_final.json'), encoding='utf-8'))
for key in ('validation_all', 'validation_indexed', 'development'):
    r = F[key]
    p, lo, hi = wilson(r['retrieved'], r['n'])
    check([round(p, 3), round(lo, 3), round(hi, 3)] == [r['share']] + r['ci95'],
          f"{key}: {r['retrieved']}/{r['n']} = {p:.1%} (95% CI {lo:.1%}-{hi:.1%})")
check(sum(x['records'] for x in F['records_term_query_by_field']) == F['records_term_query'], 'records by field add up to the unrestricted count')
check(sum(x['records'] for x in F['records_term_query_by_field'] if x['kept']) == F['records'], 'records of the kept fields add up to the reported records')

# 6. Supplementary Table S1 -------------------------------------------------------------------------
print('6. Supplementary Table S1')
wb = openpyxl.load_workbook(S1, data_only=True)
ws = wb['Data']
hdr = [c.value for c in ws[1]]
rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[hdr.index('bibkey')]]
check(len(rows) == 126 and len({r['bibkey'] for r in rows}) == 126, f'{len(rows)} included studies, unique keys')
pc = wb['PRISMA_check']
flags = [c.value for row in pc.iter_rows() for c in row if c.value in ('yes', 'no')]
check(flags and all(f == 'yes' for f in flags), f'PRISMA_check: all {len(flags)} comparisons with the flow diagram match')
split = {r['bibkey']: r['set'] for r in csv.DictReader(open(os.path.join(RC, 'split.csv'), encoding='utf-8'))}
val = {r['bibkey']: r['retrieved_by_final_query'] for r in rows if r['validation_set'] == 'validation'}
check(set(val) == {k for k, s in split.items() if s == 'validation'}, 'validation set in S1 = split.csv (38 studies)')
check(sorted(k for k, x in val.items() if x == 'no') == sorted(F['validation_missed']), 'validation studies not retrieved: S1 = results_final.json')
check(sorted(k for k, x in val.items() if x == 'not indexed in OpenAlex') == sorted(F['validation_not_indexed']), 'validation studies not indexed: S1 = results_final.json')
check(sum(x == 'yes' for x in val.values()) == F['validation_all']['retrieved'], f"validation studies retrieved: {F['validation_all']['retrieved']}")

# 7. protocol and summaries -------------------------------------------------------------------------
print('7. PROTOCOL.md and summaries')
prot = open(os.path.join(RC, 'PROTOCOL.md'), encoding='utf-8').read().split('## 9.')[1]
summ = open(os.path.join(RC, 'run_output_final', 'summary_final.txt'), encoding='utf-8').read()
v, vi = F['validation_all'], F['validation_indexed']
for label, s in (('PROTOCOL.md Section 9', prot), ('summary_final.txt', summ)):
    need = [f"{F['records']:,}", f"{v['retrieved']}/{v['n']}", f"{vi['retrieved']}/{vi['n']}", f"{100 * v['share']:.1f}%"]
    if label.startswith('PROTOCOL'):
        need += [f"{F['records_term_query']:,}", f"{100 * v['ci95'][0]:.1f}-{100 * v['ci95'][1]:.1f}%"]
    missing = [n for n in need if n not in s]
    check(not missing, f'{label}: records and retrieval as in results_final.json' + (f' (missing {missing})' if missing else ''))
    check(all(k in s for k in F['validation_missed'] + F['validation_not_indexed']), f'{label}: all validation studies not retrieved are listed')
log1 = json.load(open(os.path.join(RC, 'run_output', 'protocol_log.json'), encoding='utf-8'))
check(sha(os.path.join(RC, 'run_output', 'query_validation.txt')) == log1['query_sha256'], 'frozen term query matches its SHA-256 in run_output/protocol_log.json')

# 8. per-study outputs of the final run (written by run_final.py and count_by_source_type.py) ---------
print('8. Per-study outputs of the final run')
raw = os.path.join(RC, 'run_output_final')
need = ['retrieval_final.csv', 'protocol_log_final.json', 'records_by_field.csv', 'studies_types.csv']
for f in need:
    check(os.path.exists(os.path.join(raw, f)), f'run_output_final/{f} present')
rp = os.path.join(raw, 'retrieval_final.csv')
if os.path.exists(rp):
    rr = list(csv.DictReader(open(rp, encoding='utf-8')))
    check({r['bibkey']: r['set'] for r in rr} == split, 'retrieval_final.csv: studies and sets = split.csv')
    for st, ref in (('validation', v), ('development', F['development'])):
        check(sum(r['retrieved'] == 'True' for r in rr if r['set'] == st) == ref['retrieved'], f'retrieval_final.csv: {st} retrieved = results_final.json')
    check(sorted(r['bibkey'] for r in rr if r['set'] == 'validation' and r['retrieved'] != 'True' and r['indexed'] == 'True') == sorted(F['validation_missed']),
          'retrieval_final.csv: validation studies not retrieved = results_final.json')
lp = os.path.join(raw, 'protocol_log_final.json')
if os.path.exists(lp):
    import urllib.parse
    lf = json.load(open(lp, encoding='utf-8'))
    check(lf['records'] == F['records'] and lf['records_term_query'] == F['records_term_query'] and lf['query_sha256'] == log1['query_sha256'],
          'protocol_log_final.json: records and term query = results_final.json')
    flt = f"title.search:{urllib.parse.quote(open(os.path.join(RC, 'run_output', 'query_validation.txt'), encoding='utf-8').read().strip())},{lf['filters']}"
    check(hashlib.sha256(flt.encode('utf-8')).hexdigest() == lf['filters_sha256'], 'protocol_log_final.json: filter string matches its SHA-256')
fp = os.path.join(raw, 'records_by_field.csv')
if os.path.exists(fp):
    fr = list(csv.DictReader(open(fp, encoding='utf-8')))
    check(sum(int(r['records_term_query']) for r in fr) == F['records_term_query'], 'records_by_field.csv: all fields add up to the unrestricted count')
    check(sum(int(r['records_term_query']) for r in fr if r['kept_in_final_query'] == 'yes') == F['records'], 'records_by_field.csv: kept fields add up to the reported records')
sp = os.path.join(raw, 'studies_types.csv')
if os.path.exists(sp):
    check(len(list(csv.DictReader(open(sp, encoding='utf-8')))) == 82, 'studies_types.csv: 82 studies')
if not all(os.path.exists(os.path.join(raw, f)) for f in need):
    print('        (copy the missing files from the folder of the final run, final_run/run_output_final/)')

print('\nALL CHECKS PASSED' if all(results) else '\nSOME CHECKS FAILED')
sys.exit(0 if all(results) else 1)

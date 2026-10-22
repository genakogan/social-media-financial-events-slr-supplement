#!/usr/bin/env python3
"""Step 2 of the validated verification search: derive candidate terms from the
DEVELOPMENT titles only (see PROTOCOL.md, section 2).

Inputs : split.csv, the register (titles), the manuscript (Table 2)
Output : candidates.csv  - term, key, kind, source, df_dev, dev_categories
Validation titles are never read.
"""
import csv, re, sys, os
from collections import defaultdict
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
REGISTER = os.path.join(HERE, '..', 'Supplementary_Table_S1_Corpus_Register.xlsx')
MANUSCRIPT = os.path.join(HERE, '..', 'manuscript.tex')
OUT = os.path.join(HERE, 'candidates.csv')

STOP = set('''a an the and or of in on for to from with by via at as into onto is are was were be been being its it
their them they this that these those how what why when where which who whom whose do does did can could should would
will may might must not no nor than then there here using based toward towards versus vs new our we us your you i
between among within without over under about after before during against across through per each both all any some
more most other such only own same so too very s t just also'''.split())


def tokens(title):
    t = title.lower().replace('{', '').replace('}', '').replace('’', "'")
    t = re.sub(r"'s\b", '', t)
    return [w for w in re.findall(r'[a-z0-9]+', t) if not w.isdigit()]


def norm(w):
    """count singular and plural together (counting only)"""
    if len(w) > 5 and w.endswith('ies'):
        return w[:-3] + 'y'
    if len(w) > 4 and w.endswith('s') and not w.endswith(('ss', 'us', 'is', 'ews')):
        return w[:-1]
    return w


def grams(title):
    """set of (key, surface) for words and adjacent two-word phrases without stop words"""
    ts = tokens(title)
    out = set()
    for w in ts:
        if w not in STOP:
            out.add((norm(w), w))
    for a, b in zip(ts, ts[1:]):
        if a not in STOP and b not in STOP:
            out.add((norm(a) + ' ' + norm(b), a + ' ' + b))
    return out


def key_of(phrase):
    return ' '.join(norm(w) for w in tokens(phrase))


def contains(title, key):
    """does the title contain the (normalised) word or phrase key?"""
    seq = [norm(w) for w in tokens(title)]
    k = key.split()
    return any(seq[i:i + len(k)] == k for i in range(len(seq) - len(k) + 1))


def table2_terms(manuscript):
    m = open(manuscript, encoding='utf-8').read()
    L = m.index('\\label{tab:search_concepts}')
    tab = m[L:m.index('\\end{tabular}', L)]
    body = tab.split('\\hline', 2)[2]                      # after the header row
    terms = set()
    for row in body.split('\\\\'):
        cells = [c.strip() for c in row.split('&')]
        if len(cells) < 4:
            continue
        kw, boolean = cells[1], cells[2]
        terms |= {k.strip() for k in kw.replace('\n', ' ').split(',') if k.strip()}
        terms |= set(re.findall(r'"([^"]+)"', boolean))
        bare = re.sub(r'"[^"]+"', ' ', boolean)
        terms |= {w for w in re.findall(r'[A-Za-z][A-Za-z\-]+', bare) if w not in ('AND', 'OR')}
    return sorted(t for t in terms if t)


def main():
    split = {r['bibkey']: r for r in csv.DictReader(open(os.path.join(HERE, 'split.csv'), encoding='utf-8'))}
    ws = openpyxl.load_workbook(REGISTER, read_only=True)['Data']
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    dev = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, r))
        s = split.get(d['bibkey'])
        if s and s['set'] == 'development':          # validation titles are skipped here
            dev[d['bibkey']] = (d['title'], s['prisma_box'])
    assert len(dev) == 44, len(dev)

    df = defaultdict(set)
    surface = defaultdict(lambda: defaultdict(int))
    for key, (title, box) in dev.items():
        for k, s in grams(title):
            df[k].add(key)
            surface[k][s] += 1

    rows = {}
    for k, keys in df.items():
        if len(keys) >= 2:
            rows[k] = {'term': max(surface[k].items(), key=lambda x: (x[1], -len(x[0])))[0], 'key': k,
                       'kind': 'phrase' if ' ' in k else 'word', 'source': 'development', 'df_dev': len(keys)}
    for t in table2_terms(MANUSCRIPT):
        k = key_of(t)
        hits = [b for b, (title, _) in dev.items() if contains(title, k)]
        if not hits:
            continue
        if k in rows:
            if rows[k]['source'] == 'development':
                rows[k]['source'] = 'development; Table 2'
        else:
            rows[k] = {'term': t.lower(), 'key': k, 'kind': 'phrase' if ' ' in k else 'word', 'source': 'Table 2',
                       'df_dev': len(hits)}
    for k, r in rows.items():
        r['dev_categories'] = '; '.join(sorted({dev[b][1].split(' - ')[0] for b in dev if contains(dev[b][0], k)}))
        assert sum(contains(dev[b][0], k) for b in dev) == r['df_dev'], k

    with open(OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=['term', 'key', 'kind', 'source', 'df_dev', 'dev_categories'])
        w.writeheader()
        w.writerows(sorted(rows.values(), key=lambda r: (-r['df_dev'], r['key'])))
    print(f'{len(rows)} candidate terms -> {OUT}')


if __name__ == '__main__':
    main()

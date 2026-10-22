#!/usr/bin/env python3
"""Validated verification search (PROTOCOL.md): one command, no manual step in between.

  python3 run_validation.py --register Supplementary_Table_S1_Corpus_Register.xlsx \
      --mailto you@example.org --api-key YOUR_KEY

Steps (all mechanical, fixed in PROTOCOL.md):
  1. counts : OpenAlex title count of every candidate term (candidates.csv) and of all works (N)
  2. build  : over-representation ratio, threshold R chosen on the DEVELOPMENT titles only,
              query written to query_validation.txt and frozen (SHA-256 in protocol_log.json)
  3. run    : the frozen query is run once; records and the included studies it returns
  4. report : retrieval in the development and validation sets (Wilson 95% CI)
Outputs (folder --out, default validation_run): term_counts.csv, query_validation.txt,
  protocol_log.json, retrieval.csv, summary.txt

Needs verify_corpus.py, derive_terms.py, split.csv and candidates.csv in the same folder.
OpenAlex calls: about 90 for the counts + 1 for N + 3 for the run.
"""
import argparse, csv, datetime, hashlib, json, math, os, sys, urllib.parse

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import verify_corpus as vc          # noqa: E402  (HTTP with retries, API key, polite pool)
import derive_terms as dt           # noqa: E402  (same tokeniser / matching rule as for the candidates)
import openpyxl                     # noqa: E402

DATE = 'from_publication_date:2020-01-01,to_publication_date:2026-12-31'
TYPES = ('type:!dissertation,type:!dataset,type:!paratext,type:!editorial,type:!letter,'
         'type:!erratum,type:!peer-review,type:!retraction')
GRID = [10, 20, 50, 100, 200, 500, 1000]
DEV_TARGET = 0.90
N_DEV = 44


def q_term(term):
    return f'"{term}"' if ' ' in term else term


def count(title_query, mailto):
    flt = (f'title.search:{urllib.parse.quote(title_query)},' if title_query else '') + f'{DATE},{TYPES}'
    d = vc.get_json(f'{vc.OPENALEX}/works?filter={flt}&per-page=1&select=id', mailto)
    return (d or {}).get('meta', {}).get('count')


def wilson(k, n, z=1.96):
    if n == 0:
        return (float('nan'),) * 3
    p = k / n
    den = 1 + z * z / n
    c = (p + z * z / (2 * n)) / den
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / den
    return p, max(0.0, c - h), min(1.0, c + h)


def load_register(path):
    ws = openpyxl.load_workbook(path, read_only=True)['Data']
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    return {r[hdr.index('bibkey')]: dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)}


def build(cands, counts, N, dev_titles):
    """PROTOCOL section 3: ratio, threshold on development titles, phrase pruning. Uses development titles only."""
    for c in cands:
        n = counts[c['key']]
        c['oa_count'] = n
        c['ratio'] = math.inf if not n else (int(c['df_dev']) / N_DEV) / (n / N)

    def query_terms(R):
        kept = [c for c in cands if c['ratio'] >= R]
        words = {c['key'] for c in kept if c['kind'] == 'word'}
        return [c for c in kept if c['kind'] == 'word' or not (set(c['key'].split()) & words)]

    def dev_recall(terms):
        keys = [c['key'] for c in terms]
        return sum(any(dt.contains(t, k) for k in keys) for t in dev_titles) / len(dev_titles)

    trace = [{'R': R, 'n_terms': len(query_terms(R)), 'dev_match_offline': round(dev_recall(query_terms(R)), 3)}
             for R in GRID]
    ok = [t['R'] for t in trace if t['dev_match_offline'] >= DEV_TARGET]
    R = max(ok) if ok else min(GRID)
    terms = sorted(query_terms(R), key=lambda c: c['term'])
    return R, terms, trace


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--register', required=True)
    ap.add_argument('--mailto', required=True)
    ap.add_argument('--api-key', default=None)
    ap.add_argument('--out', default='validation_run')
    ap.add_argument('--pause', type=float, default=0.35)
    args = ap.parse_args()
    vc.API_KEY, vc.PAUSE = args.api_key, args.pause
    os.makedirs(args.out, exist_ok=True)
    log = {'protocol': 'PROTOCOL.md', 'run_started': datetime.datetime.now().isoformat(timespec='seconds'),
           'date_filter': DATE, 'type_filter': TYPES, 'grid': GRID, 'dev_target': DEV_TARGET}

    split = {r['bibkey']: r for r in csv.DictReader(open(os.path.join(HERE, 'split.csv'), encoding='utf-8'))}
    cands = list(csv.DictReader(open(os.path.join(HERE, 'candidates.csv'), encoding='utf-8')))
    reg = load_register(args.register)
    dev_titles = [reg[k]['title'] for k, s in split.items() if s['set'] == 'development']
    assert len(dev_titles) == N_DEV

    # 1. counts ---------------------------------------------------------------------------
    print(f'1/4 OpenAlex title counts for {len(cands)} candidate terms ...', flush=True)
    N = count('', args.mailto)
    counts = {}
    for i, c in enumerate(cands, 1):
        counts[c['key']] = count(q_term(c['term']), args.mailto)
        if i % 10 == 0:
            print(f'    {i}/{len(cands)}', flush=True)
    log['N_works'] = N

    # 2. build and freeze -----------------------------------------------------------------
    R, terms, trace = build(cands, counts, N, dev_titles)
    query = ' OR '.join(q_term(c['term']) for c in terms)
    with open(os.path.join(args.out, 'term_counts.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=['term', 'key', 'kind', 'source', 'df_dev', 'dev_categories', 'oa_count', 'ratio', 'in_query'])
        w.writeheader()
        qk = {c['key'] for c in terms}
        for c in sorted(cands, key=lambda c: -c['ratio']):
            w.writerow({**c, 'ratio': 'inf' if c['ratio'] == math.inf else round(c['ratio'], 2), 'in_query': c['key'] in qk})
    qpath = os.path.join(args.out, 'query_validation.txt')
    open(qpath, 'w', encoding='utf-8').write(query + '\n')
    log.update(threshold_trace=trace, R=R, n_terms=len(terms), query=query,
               query_sha256=hashlib.sha256(open(qpath, 'rb').read()).hexdigest(),
               frozen_at=datetime.datetime.now().isoformat(timespec='seconds'))
    print(f'2/4 threshold R = {R}; {len(terms)} terms; query frozen (sha256 {log["query_sha256"][:12]}...)', flush=True)

    # 3. run once -------------------------------------------------------------------------
    print('3/4 running the frozen query once ...', flush=True)
    flt = f'title.search:{urllib.parse.quote(query)},{DATE},{TYPES}'
    meta = vc.get_json(f'{vc.OPENALEX}/works?filter={flt}&per-page=1&select=id', args.mailto)
    hits = meta['meta']['count']
    ids = {k: vc.short_id(reg[k].get('openalex_id')) for k in split if reg[k].get('openalex_id')}
    found = set()
    idl = list(ids.values())
    for i in range(0, len(idl), 50):
        d = vc.get_json(f'{vc.OPENALEX}/works?filter={flt},ids.openalex:{"|".join(idl[i:i + 50])}&per-page=50&select=id',
                        args.mailto)
        found |= {vc.short_id(w['id']) for w in (d or {}).get('results', [])}

    # 4. report ---------------------------------------------------------------------------
    rows = []
    for k, s in sorted(split.items(), key=lambda x: (x[1]['set'], x[1]['prisma_box'], x[0])):
        oid = ids.get(k)
        rows.append({'bibkey': k, 'set': s['set'], 'prisma_box': s['prisma_box'], 'theme': s['theme'],
                     'openalex_id': oid or '', 'indexed': bool(oid), 'retrieved': bool(oid and oid in found),
                     'title': reg[k]['title']})
    with open(os.path.join(args.out, 'retrieval.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader(); w.writerows(rows)

    res, lines = {}, []
    for st in ('development', 'validation'):
        s_all = [r for r in rows if r['set'] == st]
        s_idx = [r for r in s_all if r['indexed']]
        k = sum(r['retrieved'] for r in s_all)
        for lab, base in (('all', s_all), ('indexed', s_idx)):
            p, lo, hi = wilson(k, len(base))
            res[f'{st}_{lab}'] = {'retrieved': k, 'n': len(base), 'share': round(p, 3), 'ci95': [round(lo, 3), round(hi, 3)]}
            lines.append(f'{st:<11} {lab:<8} {k}/{len(base)} = {p:.1%}  (95% CI {lo:.1%}-{hi:.1%})')
    missed = [f"{r['bibkey']} ({r['prisma_box']}){'' if r['indexed'] else ' - not indexed in OpenAlex'}"
              for r in rows if r['set'] == 'validation' and not r['retrieved']]
    log.update(records=hits, results=res, validation_missed=missed,
               run_finished=datetime.datetime.now().isoformat(timespec='seconds'))
    json.dump(log, open(os.path.join(args.out, 'protocol_log.json'), 'w', encoding='utf-8'), indent=2, default=str)
    summary = '\n'.join([f'records returned by the frozen query: {hits:,}', *lines, 'validation studies not retrieved:',
                         *[f'  {m}' for m in missed]])
    open(os.path.join(args.out, 'summary.txt'), 'w', encoding='utf-8').write(summary + '\n')
    print('4/4 done\n' + summary)
    print(f'\nSend the folder {args.out} back.')


if __name__ == '__main__':
    main()

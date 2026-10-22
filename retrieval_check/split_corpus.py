#!/usr/bin/env python3
"""Step 1 of the validated verification search: split the included studies published
2021-2026 into a development set and a validation set.

Rule (fixed before any term was derived):
  * population: included studies with in_search_window = yes in the register (82);
  * strata: the flow-diagram category (column prisma_box);
  * within each stratum the bibkeys are sorted, shuffled with one random.Random(SEED)
    (strata processed in sorted order), and the first ceil(n/2) go to development,
    the rest to validation.
Output: split.csv (bibkey, prisma_box, theme, set). Titles are not printed.
"""
import csv, math, random, sys, os
import openpyxl

SEED = 20261019
REGISTER = sys.argv[1] if len(sys.argv) > 1 else '../Supplementary_Table_S1_Corpus_Register.xlsx'
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'split.csv')

ws = openpyxl.load_workbook(REGISTER, read_only=True)['Data']
hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
win = [r for r in rows if r['in_search_window'] == 'yes']
assert len(win) == 82, len(win)

strata = {}
for r in win:
    strata.setdefault(r['prisma_box'], []).append(r)

rng = random.Random(SEED)
out = []
for box in sorted(strata):
    keys = sorted(r['bibkey'] for r in strata[box])
    rng.shuffle(keys)
    k = math.ceil(len(keys) / 2)
    theme = strata[box][0]['theme']
    for i, key in enumerate(keys):
        out.append({'bibkey': key, 'prisma_box': box, 'theme': theme, 'set': 'development' if i < k else 'validation'})

with open(OUT, 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=['bibkey', 'prisma_box', 'theme', 'set'])
    w.writeheader()
    w.writerows(sorted(out, key=lambda r: (r['set'], r['prisma_box'], r['bibkey'])))
n_dev = sum(r['set'] == 'development' for r in out)
print(f'seed {SEED}: {n_dev} development, {len(out) - n_dev} validation -> {OUT}')
